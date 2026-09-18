#!/usr/bin/env python3
"""Read a workbook this repository wrote, with somebody else's parsers.

    python3 scripts/check_xlsx.py

`app/lib/src/core/xlsx.dart` writes `.xlsx` by hand -- an OOXML package
is a ZIP of XML and the parts are few enough. `app/test/xlsx_test.dart`
asserts what the writer put in the ZIP, which is one implementation
asserting against itself: it can prove the bytes are what the author
intended and it cannot prove they are a workbook.

Nothing on a build machine can open Excel. What it can do is read the
file with parsers that know nothing about how it was made, and that is
what this does, twice:

  * Python's own `zipfile` and `xml.etree` for the package structure --
    every part present, every part well-formed, every sheet declared in
    the content types and reachable through a relationship;
  * and **openpyxl**, a real spreadsheet reader, for the things that
    only matter once something interprets the file: a cell's TYPE, its
    number format, its bold flag, the frozen pane, the column widths.

openpyxl is REQUIRED, not optional. A check that quietly does not run
is worse than no check -- `scripts/check_app_icons.py` gives the same
argument for decoding a PNG by hand rather than skipping when an
imaging library is absent. If the import fails, this fails.

Its WARNINGS are failures too. openpyxl said "Workbook contains no
default style" about the first version of `styles.xml`; Excel did not
mind and neither did any other reader tried, but a real reader
complaining about a file this repository generates is a signal. Turning
that warning into an error is also what caught the fix for it being
written as an XML comment containing a double hyphen, which is illegal
and left a stylesheet openpyxl refused to parse at all.

## How the sample is made

By `flutter test app/test/xlsx_sample_test.dart` with
`XLSX_SAMPLE_OUT` pointing at a temporary file. That test is a fixture
generator rather than an assertion, and it is a test rather than a
script because `report_spec.dart` imports `package:flutter/material`
for `DateTimeRange` -- so `dart run` tries to compile the Flutter
framework and stops inside its text layout code. `flutter test` is the
only runner here that executes code importing Flutter.

Built from a real `ReportSpec`, so what is checked is what the app
produces.
"""

from __future__ import annotations

import os
import pathlib
import subprocess
import sys
import tempfile
import warnings
import xml.etree.ElementTree as ET
import zipfile

MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
#: TWO relationship namespaces, and they are not interchangeable. The
#: `.rels` files use the PACKAGE one for their own elements; the `r:id`
#: attribute inside `workbook.xml` uses the OFFICEDOCUMENT one. Reading
#: `r:id` with the package namespace returns None for every sheet, which
#: is how the first version of this gate reported two sheets with no
#: relationship against a workbook openpyxl resolved perfectly.
PKG_RELS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
DOC_RELS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"

#: Every part an OOXML spreadsheet package must carry.
REQUIRED_PARTS = [
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/_rels/workbook.xml.rels",
    "xl/workbook.xml",
    "xl/styles.xml",
    "xl/worksheets/sheet1.xml",
]

problems: list[str] = []


def fail(message: str) -> None:
    problems.append(message)


def make_sample(out: pathlib.Path) -> None:
    """Run the Dart fixture generator."""
    env = dict(os.environ, XLSX_SAMPLE_OUT=str(out))
    run = subprocess.run(
        ["flutter", "test", "test/xlsx_sample_test.dart", "--reporter", "compact"],
        cwd="app",
        env=env,
        capture_output=True,
        text=True,
    )
    if run.returncode != 0:
        fail(
            "the fixture generator failed:\n"
            + "\n".join(f"      | {ln}" for ln in run.stdout.splitlines()[-25:])
            + "\n"
            + "\n".join(f"      | {ln}" for ln in run.stderr.splitlines()[-10:])
        )
        return
    if not out.exists():
        fail(
            f"{out} was not written. `xlsx_sample_test.dart` skips itself "
            "when XLSX_SAMPLE_OUT is unset -- if it skipped here, the "
            "environment did not reach `flutter test`."
        )


def check_package(path: pathlib.Path) -> None:
    """The structure, with the standard library only."""
    if not zipfile.is_zipfile(path):
        fail(f"{path} is not a ZIP, so it is not an xlsx at all")
        return

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad is not None:
            fail(f"the archive is corrupt at {bad}")
            return

        names = set(z.namelist())
        for part in REQUIRED_PARTS:
            if part not in names:
                fail(f"missing part: {part}")

        # Every part parses. A workbook whose XML is not well-formed
        # opens as a repair prompt, and that is the failure this whole
        # file exists to catch before somebody's accountant does.
        parsed: dict[str, ET.Element] = {}
        for name in sorted(names):
            if not name.endswith(".xml") and not name.endswith(".rels"):
                continue
            try:
                parsed[name] = ET.fromstring(z.read(name))
            except ET.ParseError as e:
                fail(f"{name} is not well-formed XML: {e}")
        if "xl/workbook.xml" not in parsed:
            return

        sheets = [
            (s.get("name"), s.get(f"{DOC_RELS}id"))
            for s in parsed["xl/workbook.xml"].iter(f"{MAIN}sheet")
        ]
        if len(sheets) < 2:
            fail(
                f"the sample declares {len(sheets)} sheet(s); it is built "
                "from two specs, so fewer means one was dropped"
            )

        # A sheet name Excel refuses makes the WORKBOOK unreadable, not
        # the sheet. It reports a repair prompt and says nothing about
        # which name it disliked.
        for name, _ in sheets:
            if name is None:
                fail("a sheet has no name")
                continue
            if len(name) > 31:
                fail(f"sheet name over 31 characters: {name!r}")
            for ch in r':\/?*[]':
                if ch in name:
                    fail(f"sheet name contains {ch!r}, which Excel refuses: {name!r}")
        if len({n for n, _ in sheets}) != len(sheets):
            fail(f"two sheets share a name: {[n for n, _ in sheets]}")

        # Every sheet reachable, and typed. A sheet part with no
        # Override in the content types is dropped on open with no
        # error; a sheet with no relationship cannot be found at all.
        overrides = {
            o.get("PartName")
            for o in parsed["[Content_Types].xml"].iter(f"{CT}Override")
        }
        targets = {
            r.get("Id"): r.get("Target")
            for r in parsed["xl/_rels/workbook.xml.rels"].iter(f"{PKG_RELS}Relationship")
        }
        for name, rid in sheets:
            target = targets.get(rid)
            if target is None:
                fail(f"sheet {name!r} has relationship {rid}, which is not declared")
                continue
            part = f"xl/{target}"
            if part not in names:
                fail(f"sheet {name!r} points at {part}, which is not in the archive")
            if f"/{part}" not in overrides:
                fail(
                    f"{part} has no Override in the content types, so it is "
                    "dropped silently on open"
                )

        # A `<v>` that is not a number makes the workbook damaged.
        for name, root in parsed.items():
            if not name.startswith("xl/worksheets/"):
                continue
            for cell in root.iter(f"{MAIN}c"):
                if cell.get("t") == "inlineStr":
                    continue
                v = cell.find(f"{MAIN}v")
                if v is None or v.text is None:
                    continue
                try:
                    float(v.text)
                except ValueError:
                    fail(
                        f"{name} {cell.get('r')} holds a non-numeric value in "
                        f"a numeric cell: {v.text!r}"
                    )


def check_with_openpyxl(path: pathlib.Path) -> None:
    """The semantics, with a real spreadsheet reader."""
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        fail(
            "openpyxl is not installed. It is REQUIRED: the structural "
            "check above cannot see a cell's type, its number format or "
            "the frozen pane, and a check that silently does not run is "
            "worse than no check.\n"
            "      Install it: pip install openpyxl"
        )
        return

    with warnings.catch_warnings():
        # A warning from a real reader about a file we generate is a
        # finding. See this module's docstring.
        warnings.simplefilter("error")
        try:
            book = openpyxl.load_workbook(path)
        except Exception as e:  # noqa: BLE001 - any refusal is the finding
            fail(f"openpyxl refused the workbook: {type(e).__name__}: {e}")
            return

    if "Profit & Loss" not in book.sheetnames:
        fail(
            "the ampersand in a sheet name did not survive: "
            f"{book.sheetnames}"
        )
        return
    sheet = book["Profit & Loss"]

    # THE POINT OF THE WHOLE FILE. A figure has to arrive as a NUMBER,
    # because the audit's D12 is about an accountant being able to sum a
    # column, and CSV's answer to that is a column of digits with no
    # format on it.
    amount = sheet["C6"]
    if not isinstance(amount.value, (int, float)):
        fail(
            f"C6 arrived as {type(amount.value).__name__} "
            f"({amount.value!r}); a figure written as text sums to zero"
        )
    if amount.number_format != "#,##0.00":
        fail(f"C6 has no money format: {amount.number_format!r}")

    # And an account CODE has to arrive as text, or `0100` is 100 and
    # the account cannot be found by its code.
    code = sheet["A6"]
    if code.value != "0100":
        fail(
            f"A6 arrived as {code.value!r} ({type(code.value).__name__}); "
            "an account code written as a number loses its leading nought"
        )

    # A negative stays negative. Written as a magnitude it adds up the
    # wrong way, which is the one error in a spreadsheet that looks
    # entirely correct.
    negatives = [
        c.value
        for row in sheet.iter_rows()
        for c in row
        if isinstance(c.value, (int, float)) and c.value < 0
    ]
    if not negatives:
        fail(
            "no negative figure survived. The sample carries -250.50, so "
            "either the sign was dropped or the row was"
        )

    # Bold on a total, which is the only thing making a report readable
    # once it is a grid of figures.
    if not sheet["C8"].font.bold:
        fail("the section total at C8 is not bold")

    # The heading held while the rows scroll.
    if sheet.freeze_panes != "A3":
        fail(f"the frozen pane is {sheet.freeze_panes!r}, not 'A3'")

    # The column an account name goes in. At the default width
    # "Administrative expenses — professional fees" shows as
    # "Administrativ", and widening it by hand is the first thing
    # anybody does to a CSV export.
    width = sheet.column_dimensions["B"].width
    if not width or width < 30:
        fail(f"the account-name column is {width!r} wide, which is too narrow")

    # The second sheet, whose name had a slash in it.
    if len(book.sheetnames) < 2:
        fail(f"only one sheet came back: {book.sheetnames}")
    elif "/" in book.sheetnames[1]:
        fail(f"a slash survived into a sheet name: {book.sheetnames[1]!r}")


def main() -> int:
    if not pathlib.Path("app/pubspec.yaml").exists():
        print("run this from the repository root")
        return 2

    with tempfile.TemporaryDirectory() as tmp:
        sample = pathlib.Path(tmp) / "sample.xlsx"
        make_sample(sample)
        if not problems and sample.exists():
            check_package(sample)
            check_with_openpyxl(sample)

    if problems:
        print(f"{len(problems)} problem(s) with the generated workbook:\n")
        for p in problems:
            print(f"  * {p}\n")
        return 1

    print(
        "The generated workbook reads back correctly under zipfile, "
        "xml.etree and openpyxl."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
